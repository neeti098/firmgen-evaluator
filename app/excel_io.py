from __future__ import annotations

import csv
import io
import re
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .schemas import CRITERIA


REPO_ALIASES = {
    "github", "github url", "github repo", "github repo url", "github repository",
    "github repository url", "repo", "repo url", "repository", "repository url",
    "project repository", "submission link",
}
TEAM_ALIASES = {
    "team", "team id", "team name", "project", "project name", "participant",
    "participant name", "name",
}


class SpreadsheetError(ValueError):
    pass


def _normal_header(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip().lower().replace("_", " ").replace("-", " "))


def _safe_text(value: Any) -> str:
    text = str(value or "").strip()
    if text.startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


def _find_column(headers: list[str], aliases: set[str]) -> int | None:
    for index, header in enumerate(headers):
        if header in aliases:
            return index
    for index, header in enumerate(headers):
        if any(alias in header for alias in aliases if len(alias) > 5):
            return index
    return None


def _rows_to_submissions(rows: list[list[Any]]) -> list[dict[str, str]]:
    if not rows:
        raise SpreadsheetError("The uploaded spreadsheet is empty.")
    headers = [_normal_header(value) for value in rows[0]]
    repo_index = _find_column(headers, REPO_ALIASES)
    team_index = _find_column(headers, TEAM_ALIASES)
    data_rows = rows[1:]

    if repo_index is None:
        first_row_text = [_safe_text(value) for value in rows[0]]
        if first_row_text and first_row_text[0].startswith(("http://", "https://")):
            repo_index = 0
            team_index = 1 if len(first_row_text) > 1 else None
            data_rows = rows
        else:
            raise SpreadsheetError(
                "No GitHub repository column was found. Use a header such as github_repo_url."
            )

    results: list[dict[str, str]] = []
    seen: set[tuple[str, str]] = set()
    for row_number, row in enumerate(data_rows, start=2):
        if repo_index >= len(row):
            continue
        repo_url = _safe_text(row[repo_index])
        if not repo_url:
            continue
        team_id = _safe_text(row[team_index]) if team_index is not None and team_index < len(row) else ""
        if not team_id:
            team_id = f"submission-{row_number - 1}"
        key = (team_id.lower(), repo_url.lower().rstrip("/"))
        if key in seen:
            continue
        seen.add(key)
        results.append({"team_id": team_id, "repo_url": repo_url})

    if not results:
        raise SpreadsheetError("No repository URLs were found in the uploaded spreadsheet.")
    if len(results) > 200:
        raise SpreadsheetError("A maximum of 200 submissions can be uploaded at once.")
    return results


def parse_submissions(filename: str, content: bytes) -> list[dict[str, str]]:
    lower = filename.lower()
    if lower.endswith(".csv"):
        text = content.decode("utf-8-sig", errors="replace")
        rows = [list(row) for row in csv.reader(io.StringIO(text))]
        return _rows_to_submissions(rows)
    if lower.endswith((".xlsx", ".xlsm")):
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:
            raise SpreadsheetError("The Excel file could not be read.") from exc
        sheet = workbook.active
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        workbook.close()
        return _rows_to_submissions(rows)
    raise SpreadsheetError("Upload an .xlsx or .csv file.")


def _style_header(cells) -> None:
    fill = PatternFill("solid", fgColor="14213D")
    for cell in cells:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def _autosize(sheet, minimum: int = 10, maximum: int = 44) -> None:
    for column_cells in sheet.columns:
        length = max(len(str(cell.value or "")) for cell in column_cells[:200])
        letter = get_column_letter(column_cells[0].column)
        sheet.column_dimensions[letter].width = min(max(length + 2, minimum), maximum)


def build_results_workbook(jobs: list[dict[str, Any]]) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Scores"
    summary.sheet_view.showGridLines = False

    criterion_headers = [CRITERIA[key][0] for key in CRITERIA]
    headers = [
        "Team ID", "GitHub Repository", "Project", "Status", "Commit SHA", "Total / 100",
        *criterion_headers, "Confidence", "Review Recommended", "Review Reasons",
        "Summary", "Completed / Updated",
    ]
    summary.append(headers)
    _style_header(summary[1])
    summary.freeze_panes = "A2"
    summary.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

    details = workbook.create_sheet("Evidence")
    details.sheet_view.showGridLines = False
    detail_headers = [
        "Team ID", "Project", "Criterion", "Score", "Maximum", "Confidence",
        "Rationale", "Evidence", "Missing Evidence",
    ]
    details.append(detail_headers)
    _style_header(details[1])
    details.freeze_panes = "A2"

    for job in jobs:
        result = job.get("result") or {}
        criteria_by_id = {
            item.get("criterion_id"): item for item in result.get("criteria", [])
        }
        criterion_scores = [
            criteria_by_id.get(key, {}).get("score") for key in CRITERIA
        ]
        reasons = "; ".join(result.get("review_reasons", []))
        summary.append(
            [
                _safe_text(job.get("team_id")),
                _safe_text(job.get("repo_url")),
                _safe_text(job.get("project_title")),
                job.get("status"),
                job.get("commit_sha"),
                job.get("total_score"),
                *criterion_scores,
                job.get("overall_confidence"),
                "Yes" if job.get("manual_review_recommended") else "No",
                _safe_text(reasons),
                _safe_text(result.get("project_summary", "")),
                job.get("updated_at"),
            ]
        )

        for criterion_id, (label, maximum) in CRITERIA.items():
            item = criteria_by_id.get(criterion_id)
            if not item:
                continue
            details.append(
                [
                    _safe_text(job.get("team_id")),
                    _safe_text(job.get("project_title")),
                    label,
                    item.get("score"),
                    maximum,
                    item.get("confidence"),
                    _safe_text(item.get("rationale")),
                    _safe_text("; ".join(item.get("evidence", []))),
                    _safe_text("; ".join(item.get("missing_evidence", []))),
                ]
            )

    for row in summary.iter_rows(min_row=2):
        row[5].number_format = "0.0"
        for cell in row[6:13]:
            cell.number_format = "0.0"
        row[13].number_format = "0%"
        row[14].alignment = Alignment(horizontal="center")
    for row in details.iter_rows(min_row=2):
        row[3].number_format = "0.0"
        row[4].number_format = "0.0"
        row[5].number_format = "0%"

    thin = Side(style="thin", color="D8DEE9")
    for sheet in (summary, details):
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = Border(bottom=thin)
        sheet.row_dimensions[1].height = 34

    _autosize(summary)
    _autosize(details)
    summary.column_dimensions["B"].width = 42
    summary.column_dimensions["Q"].width = 42
    summary.column_dimensions["R"].width = 55
    details.column_dimensions["G"].width = 60
    details.column_dimensions["H"].width = 48
    details.column_dimensions["I"].width = 48

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
